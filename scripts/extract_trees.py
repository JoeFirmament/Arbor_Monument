#!/usr/bin/env python3
"""Extract Suzhou ancient-tree records from the official Word catalogues."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "source-docs"

SOURCES = {
    "张家港市": {
        "file": "张家港市.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/89d6aa2394994bb4b4579dedf4858c90.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/89d6aa2394994bb4b4579dedf4858c90/files/97315108081f4d83bb479308a7c3b831.docx",
        "center": (31.8650, 120.5550),
        "spread": (0.105, 0.145),
    },
    "常熟市": {
        "file": "常熟市.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/b942e632fd594c62ad810da5b62fe35c.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/b942e632fd594c62ad810da5b62fe35c/files/17f5ab103c8e482997442593852b7abb.docx",
        "center": (31.6544, 120.7520),
        "spread": (0.115, 0.135),
    },
    "太仓市": {
        "file": "太仓市.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/d5bc941f6d9d4dbc9dd4c404c71e190e.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/d5bc941f6d9d4dbc9dd4c404c71e190e/files/71bb89e5f00b4854af65ec5c400d9c8b.docx",
        "center": (31.4590, 121.1290),
        "spread": (0.095, 0.085),
    },
    "昆山市": {
        "file": "昆山市.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/0838e48f2ab44181a0bdc6bffcbe616f.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/0838e48f2ab44181a0bdc6bffcbe616f/files/5684e08ec2f241519675874fea4132c2.docx",
        "center": (31.3850, 120.9810),
        "spread": (0.095, 0.125),
    },
    "吴江区": {
        "file": "吴江区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/76c859abcaff4743a7f19c0d8d1d1eaf.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/76c859abcaff4743a7f19c0d8d1d1eaf/files/caeb98a84023445498b4b377d2e027e8.docx",
        "center": (31.1590, 120.6410),
        "spread": (0.155, 0.115),
    },
    "吴中区": {
        "file": "吴中区.doc",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/149aa049258948e4b6588ba53e58247c.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/149aa049258948e4b6588ba53e58247c/files/415d47168e944733809e7774f3aaa20f.doc",
        "center": (31.1880, 120.4850),
        "spread": (0.135, 0.185),
    },
    "相城区": {
        "file": "相城区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/1209c3a99b8748988485de7d3e48b78d.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/1209c3a99b8748988485de7d3e48b78d/files/bcff83c742bc4a75a4e3e1957f6a9579.docx",
        "center": (31.4390, 120.6420),
        "spread": (0.080, 0.095),
    },
    "姑苏区": {
        "file": "姑苏区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/c39761b4712441d9af389bf1cf5b1098.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/c39761b4712441d9af389bf1cf5b1098/files/25574bf72502405aa8e3fff2c781aedf.docx",
        "center": (31.3150, 120.6190),
        "spread": (0.036, 0.042),
    },
    "工业园区": {
        "file": "工业园区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/7989ff98b1bf451e8b8367c7748646cb.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/7989ff98b1bf451e8b8367c7748646cb/files/d3b4077f9f3b4529b536a0b24359bf02.docx",
        "center": (31.3240, 120.7300),
        "spread": (0.047, 0.065),
    },
    "高新区": {
        "file": "高新区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/3ae53bce4cb84b6da147df23390cae7f.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/3ae53bce4cb84b6da147df23390cae7f/files/cc508223b25d4f0d9f5baf12bb5bd7d7.docx",
        "center": (31.3050, 120.4470),
        "spread": (0.070, 0.095),
    },
    "园林景区": {
        "file": "园林景区.docx",
        "page": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/7896e4cdc7574224ac5957f55d661662.shtml",
        "document": "https://ylj.suzhou.gov.cn/szsylj/zygk/202502/7896e4cdc7574224ac5957f55d661662/files/233375cfaaea4ccc876152563d29f991.docx",
        "center": (31.3220, 120.6020),
        "spread": (0.045, 0.055),
    },
}

# Corrections verified against authoritative botanical databases. The official
# tables occasionally collapse spaces when Word cells are extracted.
SCIENTIFIC_NAME_CORRECTIONS = {
    "牛鼻栓": "Fortunearia sinensis Rehder & E.H.Wilson",
}


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def number(value: str) -> float | int | None:
    match = re.search(r"-?\d+(?:\.\d+)?", clean(value))
    if not match:
        return None
    parsed = float(match.group())
    return int(parsed) if parsed.is_integer() else parsed


def grade(age: int | float | None) -> str:
    if age is None:
        return "未标注"
    if age >= 500:
        return "一级"
    if age >= 300:
        return "二级"
    return "三级"


def approximate_point(region: str, key: str) -> tuple[float, float]:
    """Deterministic district-level display point, explicitly not a geocode."""
    source = SOURCES[region]
    lat, lng = source["center"]
    lat_spread, lng_spread = source["spread"]
    digest = hashlib.sha256(key.encode("utf-8")).digest()
    angle = int.from_bytes(digest[:4], "big") / 2**32 * math.tau
    radius = math.sqrt(int.from_bytes(digest[4:8], "big") / 2**32) * 0.92
    return (
        round(lat + math.sin(angle) * lat_spread * radius, 6),
        round(lng + math.cos(angle) * lng_spread * radius, 6),
    )


def convert_legacy_doc(path: Path, out_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("soffice is required to read the legacy .doc catalogue")
    profile = out_dir / "libreoffice-profile"
    profile.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "docx",
            "--outdir",
            str(out_dir),
            str(path),
        ],
        check=True,
        capture_output=True,
    )
    converted = out_dir / f"{path.stem}.docx"
    if not converted.exists():
        raise RuntimeError(f"Failed to convert {path.name}")
    return converted


def main() -> None:
    records: list[dict] = []
    with tempfile.TemporaryDirectory(prefix="suzhou-trees-") as temp:
        temp_dir = Path(temp)
        for region, source in SOURCES.items():
            source_path = SOURCE_DIR / source["file"]
            docx_path = (
                convert_legacy_doc(source_path, temp_dir)
                if source_path.suffix.lower() == ".doc"
                else source_path
            )
            document = Document(docx_path)
            if not document.tables:
                raise RuntimeError(f"No table found in {source_path.name}")
            table = document.tables[0]
            for row in table.rows[1:]:
                cells = [clean(cell.text) for cell in row.cells]
                if len(cells) < 8 or not cells[0] or cells[0].startswith("备注"):
                    continue
                tree_age = number(cells[3])
                uid = f"{region}-{cells[0]}"
                lat, lng = approximate_point(region, f"{uid}-{cells[7]}")
                records.append(
                    {
                        "uid": uid,
                        "catalog": region,
                        "number": cells[0],
                        "species": cells[1],
                        "scientificName": SCIENTIFIC_NAME_CORRECTIONS.get(cells[1], cells[2]),
                        "age": tree_age,
                        "heightM": number(cells[4]),
                        "girthCm": number(cells[5]),
                        "canopyM": number(cells[6]),
                        "address": cells[7],
                        "grade": grade(tree_age),
                        "lat": lat,
                        "lng": lng,
                        "coordinateQuality": "地区级示意",
                        "sourcePage": source["page"],
                        "sourceDocument": source["document"],
                    }
                )

    if len(records) != 2307:
        raise RuntimeError(f"Expected 2307 records, extracted {len(records)}")

    data_dir = ROOT / "data"
    public_dir = ROOT / "public" / "data"
    public_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "meta": {
            "title": "苏州市古树名木录",
            "recordCount": len(records),
            "catalogCount": len(SOURCES),
            "sourcePublished": "2025-02",
            "surveyYear": 2024,
            "coordinateNote": "官方附件未提供经纬度。地图点位按名录所属地区做确定性分布示意，不代表树木精确位置；具体位置以原始地址字段为准。",
            "sourceIndex": "https://ylj.suzhou.gov.cn/szsylj/szgsmm/szgsmzztzl.shtml",
        },
        "records": records,
    }
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    (data_dir / "trees.json").write_text(serialized, encoding="utf-8")
    (public_dir / "trees.json").write_text(serialized, encoding="utf-8")

    fields = [
        "uid",
        "catalog",
        "number",
        "species",
        "scientificName",
        "age",
        "heightM",
        "girthCm",
        "canopyM",
        "address",
        "grade",
        "lat",
        "lng",
        "coordinateQuality",
        "sourcePage",
        "sourceDocument",
    ]
    for csv_path in (data_dir / "trees.csv", public_dir / "trees.csv"):
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(records)

    print(f"Extracted {len(records)} records from {len(SOURCES)} official catalogues")


if __name__ == "__main__":
    main()
